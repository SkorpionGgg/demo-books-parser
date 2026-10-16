from bs4 import BeautifulSoup

import books_parser as bp

PAGE = """
<div class="side_categories"><ul><li><a href="catalogue/category/books_1/index.html">Books</a>
  <ul>
    <li><a href="catalogue/category/books/travel_2/index.html"> Travel </a></li>
    <li><a href="catalogue/category/books/science-fiction_16/index.html">
        Science Fiction</a></li>
  </ul></li></ul></div>
<article class="product_pod">
  <p class="star-rating Three"></p>
  <h3><a href="a-light-in-the-attic_1000/index.html" title="A Light in the Attic">A Light in the ...</a></h3>
  <div class="product_price"><p class="price_color">£51.77</p>
  <p class="instock availability"> In stock </p></div>
</article>
<article class="product_pod">
  <p class="star-rating Five"></p>
  <h3><a href="sold-out_1/index.html" title="Sold Out">Sold Out</a></h3>
  <p class="price_color">Â£1,020.50</p>
  <p class="availability">Out of stock</p>
</article>
<ul class="pager"><li class="next"><a href="page-2.html">next</a></li></ul>
"""

URL = "https://books.toscrape.com/catalogue/page-1.html"


def soup():
    return BeautifulSoup(PAGE, "html.parser")


def test_books_parsed():
    books = bp.parse_books(soup(), URL)
    assert len(books) == 2
    first, second = books
    assert first.title == "A Light in the Attic"
    assert (first.price, first.currency, first.rating, first.in_stock) == (51.77, "GBP", 3, True)
    assert first.url == "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html"
    assert (second.price, second.rating, second.in_stock) == (1020.5, 5, False)


def test_next_page():
    assert bp.next_page(soup(), URL) == "https://books.toscrape.com/catalogue/page-2.html"
    assert bp.next_page(BeautifulSoup("<ul></ul>", "html.parser"), URL) is None


def test_categories():
    cats = bp.parse_categories(soup())
    assert set(cats) == {"travel", "science fiction"}
    assert cats["travel"].endswith("/travel_2/index.html")


def test_save_files(tmp_path):
    books = bp.parse_books(soup(), URL)
    bp.save_csv(books, tmp_path / "b.csv")
    bp.save_xlsx(books, tmp_path / "b.xlsx")
    text = (tmp_path / "b.csv").read_text(encoding="utf-8-sig").splitlines()
    assert text[0] == "Название;Цена;Валюта;Рейтинг;В наличии;Ссылка"
    assert len(text) == 3
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "b.xlsx").active
    assert ws.max_row == 3 and ws["A2"].value == "A Light in the Attic"
    assert ws["F2"].hyperlink.target.startswith("https://")


def test_args_validation():
    args = bp.build_args(["-c", "travel", "-p", "2", "-o", "out/x"])
    assert (args.category, args.pages, args.output) == ("travel", 2, "out/x")
